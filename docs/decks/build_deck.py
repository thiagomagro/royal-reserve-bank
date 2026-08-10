from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
NAVY = RGBColor(0x0B, 0x25, 0x45)
BODY = RGBColor(0x2E, 0x34, 0x40)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
MID = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Liberation Sans"


def set_run(run, size, color=BODY, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_box(slide, x, y, w, h, fill=None, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or WHITE
    shape.line.color.rgb = line or fill or WHITE
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_text(slide, text, x, y, w, h, size=19, color=BODY, bold=False,
             italic=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, color, bold, italic)
    return box


def add_bullets(slide, bullets, x=0.8, y=2.1, w=11.75, h=3.7, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {text}"
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = BODY
        p.space_after = Pt(18)
        p.line_spacing = 1.12
    return box


def add_title(slide, title, number):
    add_text(slide, f"{number:02d}", 0.55, 0.43, 0.45, 0.3, 12, GOLD, True)
    add_text(slide, title, 1.05, 0.27, 11.65, 0.88, 30, NAVY, True)
    add_box(slide, 1.05, 1.34, 11.7, 0.025, fill=GOLD)


def add_footer(slide, text, y=6.73):
    add_text(slide, text, 0.72, y, 11.9, 0.38, 12, MID, italic=True, margin=0)


def add_standard_slide(prs, number, title, bullets, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, number)
    add_bullets(slide, bullets)
    add_footer(slide, note)
    return slide


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1 — cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(slide, 0, 0, 13.333, 0.18, fill=GOLD)
    add_text(slide, "RISK SCORE POR CLIENTE", 0.82, 0.72, 5.5, 0.34, 13, GOLD, True)
    add_text(slide, "Hoje o banco opera sem nenhuma medida de risco por cliente",
             0.8, 1.45, 11.7, 1.25, 31, NAVY, True)
    add_box(slide, 0.82, 2.95, 1.15, 0.035, fill=GOLD)
    add_bullets(slide, [
        "Contas e transações são tratadas sem diferenciação de perfil.",
        "O objetivo é um risk score por cliente, consultável em tempo real e recalculado diariamente.",
    ], x=0.82, y=3.35, w=11.75, h=2.15, size=20)
    add_footer(slide, "Comitê executivo · Decisão pedida: aprovar a Abordagem B e a Fase 0 · Fonte: design doc Risk Score por Cliente", y=6.68)

    add_standard_slide(
        prs, 2,
        "O bloqueio não é o cálculo do score, é a ausência dos dados",
        [
            'Não existe entidade "cliente": a conta guarda apenas nome do titular, saldo e moeda, e o nome não é único.',
            "A transação não registra dono, valor monetário nem data, então não há histórico transacional por cliente.",
        ],
        "Nenhum dos cinco fatores pedidos (saldo, volume, frequência, tempo de casa, KYC) existe hoje.",
    )
    add_standard_slide(
        prs, 3,
        "Uma Fase 0 de modelo de dados é pré-requisito de qualquer abordagem",
        [
            "Criar a entidade Customer no account-api, com documento, data de abertura e KYC, e ligar as contas a ela.",
            "Incluir cliente, valor com moeda e timestamp na transação e no evento publicado no Kafka.",
        ],
        "Sem a Fase 0, só é possível entregar um score de fachada.",
    )

    # Slide 4 — comparison table
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Três abordagens, com o mesmo pré-requisito e coberturas diferentes", 4)
    rows, cols = 4, 5
    data = [
        ["Abordagem", "Atende realtime 300 ms", "Isolamento de falha", "Esforço", "Confiança"],
        ["A — score dentro do account-api", "Só se materializar o score", "Baixo", "M", "75"],
        ["B — novo risk-score-api com eventos e job", "Sim", "Alto", "G", "70"],
        ["C — risk-score-api só em batch", "Leitura rápida, dado defasado", "Alto", "M", "80"],
    ]
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.65), Inches(1.58), Inches(12.03), Inches(3.72))
    table = table_shape.table
    widths = [3.25, 3.0, 2.0, 1.15, 2.63]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.08)
            cell.margin_bottom = Inches(0.06)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 else WHITE)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c < 3 else PP_ALIGN.CENTER
                for run in p.runs:
                    set_run(run, 13 if r == 0 else 12.5, WHITE if r == 0 else BODY, r == 0)
    add_footer(slide, "Confiança de sucesso da implantação conforme o design doc.", y=6.65)

    add_standard_slide(
        prs, 5,
        "A recomendação é a Abordagem B, precedida da Fase 0",
        [
            "O SLA de 300 ms com recálculo diário exige score materializado e cálculo fora do caminho de leitura.",
            "Colocar o motor de risco dentro do account-api acopla risco a cadastro sem economizar o custo da persistência.",
        ],
        "A Abordagem C só se justifica se o negócio aceitar score defasado entre execuções do job.",
    )
    add_standard_slide(
        prs, 6,
        "O impacto é medido pela cobertura dos fatores e pela auditabilidade do score",
        [
            "Score 0–1000 (BAIXO 0–299, MÉDIO 300–599, ALTO 600–1000) com quatro componentes: KYC 30%, tempo de casa 20%, saldo agregado em BRL 20%, comportamento transacional 30%.",
            "Pesos e limiares parametrizáveis via Config Server e fatores persistidos a cada cálculo, para recalibrar sem deploy e sustentar auditoria.",
        ],
        "Pendência de dado: régua depende de validação escrita de Risco/Crédito; não há meta de negócio quantificada na fonte.",
    )
    add_standard_slide(
        prs, 7,
        "Os riscos maiores são de dado, de acesso e de compliance, e todos têm mitigação definida",
        [
            "Sem role de admin no provedor de identidade, o serviço não publica rota no gateway e opera apenas serviço-a-serviço.",
            "LGPD e discriminação algorítmica exigem validação do Jurídico antes de persistir score e operação em modo shadow antes de qualquer decisão de crédito.",
        ],
        "Câmbio via boletim de fechamento do BCB, com cache da última cotação e registro da taxa usada.",
    )
    add_standard_slide(
        prs, 8,
        "Pedimos a aprovação da Abordagem B com a Fase 0 como épico priorizado",
        [
            "Decisão pedida: aprovar a Abordagem B e priorizar a Fase 0 como épico próprio, anterior ao score.",
            "Próximos passos: Fase 0, risk-score-api, job diário, modo shadow, role de admin e publicação da rota.",
        ],
        "Perguntas em aberto (do design doc): role de admin; fatores econômicos internos; unificação de titulares legados; confirmação escrita dos valores neutros de cliente legado e KYC ausente; qual boletim do BCB é o oficial.",
    )

    prs.save(OUT / "risk-score-deck.pptx")


if __name__ == "__main__":
    build()
